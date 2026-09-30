"""Excel workbook for consolidated RF Check results."""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from reports.services.payload import ExportPayload, ValidationRow

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="2B6CB0")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_TITLE_FONT = Font(bold=True, color="1A1A18", name="Calibri", size=14)
_LABEL_FONT = Font(bold=True, color="1A1A18", name="Calibri", size=11)
_BODY_FONT = Font(color="1A1A18", name="Calibri", size=11)
_OK_FILL = PatternFill("solid", fgColor="DCFCE7")
_WARN_FILL = PatternFill("solid", fgColor="FEF3C7")
_NOK_FILL = PatternFill("solid", fgColor="FEE2E2")
_THIN = Border(
    left=Side(style="thin", color="E7E5E4"),
    right=Side(style="thin", color="E7E5E4"),
    top=Side(style="thin", color="E7E5E4"),
    bottom=Side(style="thin", color="E7E5E4"),
)
_WRAP = Alignment(wrap_text=True, vertical="center")

_VALIDATION_HEADERS = (
    "Technology",
    "Site",
    "Check type",
    "Parameter",
    "Expected Value",
    "Found Value",
    "Status",
    "Notes",
)


def render_xlsx(payload: ExportPayload) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, payload)
    _write_validations(workbook.create_sheet("Validations"), payload.validations)
    _write_validations(workbook.create_sheet("OK"), payload.rows_by_status("consistent"))
    _write_validations(workbook.create_sheet("Warnings"), payload.rows_by_status("failed"))
    _write_validations(workbook.create_sheet("NOK"), payload.rows_by_status("inconsistent"))

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def report_basename(payload: ExportPayload) -> str:
    site = re.sub(r"[^A-Za-z0-9._-]+", "_", payload.site).strip("._-") or "site"
    tech = re.sub(r"[^A-Za-z0-9]+", "", payload.technology) or "RF"
    return f"Claro_RF_Check_Report_{tech}_{site}"


def _write_summary(sheet: Worksheet, payload: ExportPayload) -> None:
    sheet["A1"] = "Claro RF Check — Analysis report"
    sheet["A1"].font = _TITLE_FONT
    sheet.merge_cells("A1:B1")

    rows = [
        ("Status", payload.analysis_status_label),
        ("Technology", payload.tech_label),
        ("Site", payload.site),
        ("Type", payload.check_types_label),
        ("Date", payload.processed_at),
        ("Cells", str(len(payload.cells))),
        ("Selected cells", ", ".join(payload.cells) or "—"),
        ("Overall Score", f"{payload.score}%" if payload.score is not None else "—"),
        ("Total", payload.counts["total"]),
        ("OK", payload.counts["consistent"]),
        ("Warnings", payload.counts["failed"]),
        ("NOK", payload.counts["inconsistent"]),
        ("EP file", payload.ep_filename or "—"),
        ("Pre-check file", payload.precheck_filename or "—"),
        ("Full Check file", payload.full_check_filename or "—"),
        ("Analysis ID", payload.analysis_id),
    ]
    for index, (label, value) in enumerate(rows, start=3):
        sheet.cell(index, 1, label).font = _LABEL_FONT
        sheet.cell(index, 2, value).font = _BODY_FONT

    start = 3 + len(rows) + 1
    sheet.cell(start, 1, "Results by check type").font = _TITLE_FONT
    headers = ("Check type", "Status", "Date", "Total", "OK", "Warnings", "NOK", "Score")
    header_row = start + 1
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, col, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        cell.border = _THIN
    for offset, section in enumerate(payload.sections):
        values = (
            section.check_type,
            section.status_label,
            section.processed_at,
            section.counts["total"],
            section.counts["consistent"],
            section.counts["failed"],
            section.counts["inconsistent"],
            f"{section.score}%" if section.score is not None else "—",
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(header_row + 1 + offset, col, value)
            cell.font = _BODY_FONT
            cell.alignment = _WRAP
            cell.border = _THIN

    _autosize(sheet, max_col=8)


def _write_validations(sheet: Worksheet, rows: list[ValidationRow]) -> None:
    for col, header in enumerate(_VALIDATION_HEADERS, start=1):
        cell = sheet.cell(1, col, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        cell.border = _THIN
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{max(len(rows) + 1, 1)}"

    for index, row in enumerate(rows, start=2):
        values = (
            row.technology,
            row.site,
            row.check_type,
            row.label,
            row.expected,
            row.found,
            row.status_label,
            row.note,
        )
        fill = _status_fill(row.status)
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(index, col, value)
            cell.font = _BODY_FONT
            cell.alignment = _WRAP
            cell.border = _THIN
            if fill is not None:
                cell.fill = fill
    _autosize(sheet, max_col=len(_VALIDATION_HEADERS))


def _status_fill(status: str) -> PatternFill | None:
    if status == "consistent":
        return _OK_FILL
    if status == "failed":
        return _WARN_FILL
    if status == "inconsistent":
        return _NOK_FILL
    return None


def _autosize(sheet: Worksheet, *, max_col: int) -> None:
    for col in range(1, max_col + 1):
        width = 12
        for cell in sheet[get_column_letter(col)]:
            if cell.value is None:
                continue
            width = max(width, min(len(str(cell.value)) + 2, 48))
        sheet.column_dimensions[get_column_letter(col)].width = width
