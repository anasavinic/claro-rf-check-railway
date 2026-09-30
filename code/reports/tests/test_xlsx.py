from io import BytesIO

import pytest
from openpyxl import load_workbook

from precheck.models import CheckAnalysisStatus
from reports.services.payload import build_export_payload
from reports.services.xlsx import render_xlsx, report_basename
from reports.tests.conftest import make_analysis, make_precheck_result


@pytest.mark.django_db
def test_xlsx_contains_same_rows_as_ui(ep_job, user):
    analysis = make_analysis(ep_job, user, status=CheckAnalysisStatus.INCONSISTENT)
    make_precheck_result(analysis)
    payload = build_export_payload(analysis)

    workbook = load_workbook(BytesIO(render_xlsx(payload)))

    assert workbook.sheetnames == ["Summary", "Validations", "OK", "Warnings", "NOK"]
    summary = workbook["Summary"]
    assert summary["B3"].value == "Inconsistent"
    assert summary["B4"].value == "5G NR"
    assert summary["B5"].value == "S01PRCLG21"
    assert summary["B10"].value == "50%"

    validations = workbook["Validations"]
    assert [cell.value for cell in validations[1]] == [
        "Technology",
        "Site",
        "Check type",
        "Parameter",
        "Expected Value",
        "Found Value",
        "Status",
        "Notes",
    ]
    assert [cell.value for cell in validations[2][3:8]] == [
        "gNodeB ID",
        "1060541",
        "1060541",
        "OK",
        "Matches EP gNBId.",
    ]
    assert [cell.value for cell in validations[3][3:8]] == [
        "Tracking Area Code",
        "4151041",
        "4151040",
        "NOK",
        "Network Tracking Area Code differs from EP Tracking Area ID.",
    ]
    assert workbook["OK"].max_row == 2
    assert workbook["NOK"].max_row == 2
    assert workbook["Warnings"].max_row == 1
    assert report_basename(payload) == "Claro_RF_Check_Report_5G_S01PRCLG21"
