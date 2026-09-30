"""Orchestrates EP Claro import: validate package, then stream cells."""

from __future__ import annotations

from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from ep_import.schema import CELL_KEY_COLUMNS, MAX_ROW_ISSUES, SHEET_BY_TECH, Technology, region_from_filename
from ep_import.services.content import validate_file_meta, validate_xlsx_package
from ep_import.services.errors import (
    EpImportError,
    EpImportIssue,
    cap_issues,
    duplicate_cell,
    empty_sheet,
    invalid_workbook,
)
from ep_import.services.layout import validate_headers, validate_sheets_present
from ep_import.services.mapper import build_header_map, canonicalize_row
from ep_import.services.normalizer import ImportPayload, NormalizedCell, inspect_row
from ep_import.services.reader import list_extra_sheets, sheet_rows
from ep_import.services.source import open_workbook_handle, owned_workbook_stream, source_name, source_size

ProgressFn = Callable[[str, int, int | None], None]
CellKey = tuple[str, str, str]
SeenCells = dict[CellKey, tuple[str, int]]


@dataclass
class ValidationResult:
    filename: str
    region: str | None
    warnings: list[EpImportIssue] = field(default_factory=list)
    ignored_sheets: list[str] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=dict)
    row_count: int = 0


def import_ep_file(
    source: str | Path | BinaryIO,
    filename: str | None = None,
    size: int | None = None,
    progress: ProgressFn | None = None,
) -> ImportPayload:
    """Import a Claro EP workbook. Raises EpImportError when layout or rows are invalid."""
    name = source_name(source, filename)
    with owned_workbook_stream(source) as stream:
        result = validate_workbook(stream, name, size=source_size(source, size), progress=progress)
        stream.seek(0)
        cells = list(iter_valid_cells(stream, name, file_region=result.region))
    return _payload_from(result, cells)


def validate_workbook(
    stream: BinaryIO,
    filename: str,
    size: int | None = None,
    progress: ProgressFn | None = None,
) -> ValidationResult:
    """Validate package, headers, and keys. Duplicate cells become warnings (first wins)."""
    _notify(progress, "validating", 0, None)
    meta_issues = validate_file_meta(filename, size)
    if meta_issues:
        raise EpImportError(title="Invalid EP file", issues=meta_issues)

    content_issues = validate_xlsx_package(stream, filename)
    if content_issues:
        raise EpImportError(title="Invalid EP file", issues=content_issues)

    file_region = region_from_filename(filename)
    region = file_region.value if file_region else None
    seen: SeenCells = {}
    row_issues: list[EpImportIssue] = []
    warnings: list[EpImportIssue] = []
    counts = {tech.value: 0 for tech in Technology}
    processed = 0
    ignored: list[str] = []

    try:
        with open_workbook_handle(stream) as handle:
            workbook = load_workbook(handle, read_only=True, data_only=True)
            try:
                missing = validate_sheets_present(list(workbook.sheetnames))
                if missing:
                    raise EpImportError(
                        title="Layout is incompatible with Claro SRAN EP",
                        issues=missing,
                    )
                ignored = list_extra_sheets(list(workbook.sheetnames))
                layout_issues: list[EpImportIssue] = []
                for tech, sheet_name in SHEET_BY_TECH.items():
                    headers, records = sheet_rows(workbook[sheet_name])
                    header_issues = validate_headers(sheet_name, headers, tech)
                    if header_issues:
                        layout_issues.extend(header_issues)
                        for _record in records:
                            pass
                        continue
                    header_map = build_header_map(headers)
                    sheet_rows_seen = 0
                    for excel_row, record in records:
                        processed += 1
                        sheet_rows_seen += 1
                        if processed % 500 == 0:
                            _notify(progress, "validating", processed, None)
                        cell, issues = inspect_row(
                            canonicalize_row(record, header_map),
                            tech,
                            sheet_name,
                            excel_row,
                            region,
                        )
                        if issues:
                            row_issues.extend(issues)
                            continue
                        assert cell is not None
                        if not region and cell.region:
                            region = cell.region
                        dup = _duplicate_warning(seen, cell, sheet_name, excel_row)
                        if dup:
                            warnings.append(dup)
                            continue
                        counts[cell.technology.value] += 1
                    if sheet_rows_seen == 0:
                        warnings.append(empty_sheet(sheet_name))
                if layout_issues:
                    raise EpImportError(
                        title="Layout is incompatible with Claro SRAN EP",
                        issues=cap_issues(layout_issues + row_issues, MAX_ROW_ISSUES),
                    )
            finally:
                workbook.close()
    except EpImportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise EpImportError(title="Could not open EP file", issues=[invalid_workbook()]) from exc

    if row_issues:
        raise EpImportError(
            title="EP file has invalid rows",
            issues=cap_issues(row_issues, MAX_ROW_ISSUES),
        )

    _notify(progress, "validating", processed, processed)
    return ValidationResult(
        filename=filename,
        region=region,
        warnings=cap_issues(warnings, MAX_ROW_ISSUES),
        ignored_sheets=ignored,
        counts=counts,
        row_count=processed,
    )


def iter_valid_cells(
    stream: BinaryIO,
    filename: str,
    file_region: str | None = None,
) -> Iterator[NormalizedCell]:
    """Second pass. Caller must have validated the stream already. First occurrence wins."""
    region = file_region
    seen: SeenCells = {}
    with open_workbook_handle(stream) as handle:
        workbook = load_workbook(handle, read_only=True, data_only=True)
        try:
            for tech, sheet_name in SHEET_BY_TECH.items():
                headers, records = sheet_rows(workbook[sheet_name])
                header_map = build_header_map(headers)
                for excel_row, record in records:
                    cell, issues = inspect_row(
                        canonicalize_row(record, header_map),
                        tech,
                        sheet_name,
                        excel_row,
                        region,
                    )
                    if issues or cell is None:
                        continue
                    if not region and cell.region:
                        region = cell.region
                    if _duplicate_warning(seen, cell, sheet_name, excel_row):
                        continue
                    yield cell
        finally:
            workbook.close()


def _duplicate_warning(
    seen: SeenCells,
    cell: NormalizedCell,
    sheet_name: str,
    excel_row: int,
) -> EpImportIssue | None:
    """Register cell key; return a warning if this key was already seen (first wins)."""
    key: CellKey = (cell.technology.value, cell.site_name, cell.cell_name)
    previous = seen.get(key)
    if previous:
        return duplicate_cell(
            sheet=sheet_name,
            row=excel_row,
            column=CELL_KEY_COLUMNS[cell.technology],
            technology=cell.technology.value,
            site=cell.site_name,
            cell=cell.cell_name,
            first_sheet=previous[0],
            first_row=previous[1],
        )
    seen[key] = (sheet_name, excel_row)
    return None


def _payload_from(result: ValidationResult, cells: list[NormalizedCell]) -> ImportPayload:
    payload = ImportPayload(
        filename=result.filename,
        region=result.region,
        cells=cells,
        warnings=result.warnings,
        ignored_sheets=result.ignored_sheets,
    )
    return payload


def _notify(progress: ProgressFn | None, stage: str, processed: int, total: int | None) -> None:
    if progress is None:
        return
    progress(stage, processed, total)
