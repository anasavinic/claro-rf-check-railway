"""XLSX reader for Claro EP — canonical RNP sheets, one row at a time."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from ep_import.schema import REQUIRED_SHEETS
from ep_import.services.errors import EpImportError, invalid_workbook
from ep_import.services.source import open_workbook_handle


def read_sheet_headers(source: str | Path | BinaryIO) -> dict[str, list[str]]:
    """Return {sheet_name: headers} for all sheets (including extras)."""
    with _open_workbook(source) as wb:
        return _headers_from_workbook(wb)


def read_canonical_sheets(source: str | Path | BinaryIO) -> dict[str, list[dict[str, Any]]]:
    """Materialize required sheets. Used by tests; the job path streams instead."""
    data: dict[str, list[dict[str, Any]]] = {}
    with _open_workbook(source) as wb:
        missing = [name for name in sorted(REQUIRED_SHEETS) if name not in wb.sheetnames]
        if missing:
            from ep_import.services.errors import missing_sheet

            raise EpImportError(
                title="Layout is incompatible with Claro SRAN EP",
                issues=[missing_sheet(name) for name in missing],
            )
        for sheet in REQUIRED_SHEETS:
            _headers, records = sheet_rows(wb[sheet])
            data[sheet] = [record for _row, record in records]
    return data


def sheet_rows(ws) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    """One forward pass: header row, then (excel_row, record) for non-empty data rows."""
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    headers = _clean_headers(header_row or ())
    return headers, _data_records(rows_iter, headers)


def iter_sheet_records(ws) -> Iterator[tuple[int, dict[str, Any]]]:
    """Yield (excel_row, record) for non-empty data rows. Does not keep the sheet."""
    _headers, records = sheet_rows(ws)
    return records


def _data_records(rows_iter, headers: list[str]) -> Iterator[tuple[int, dict[str, Any]]]:
    for excel_row, row in enumerate(rows_iter, start=2):
        if row is None or _is_empty_row(row):
            continue
        record: dict[str, Any] = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = row[idx] if idx < len(row) else None
            record[header] = _json_safe_value(value)
        if any(value not in (None, "") for value in record.values()):
            yield excel_row, record


def list_extra_sheets(sheet_names: list[str]) -> list[str]:
    return [name for name in sheet_names if name not in REQUIRED_SHEETS]


def sheet_header_map(wb: Workbook) -> dict[str, list[str]]:
    return _headers_from_workbook(wb)


def _headers_from_workbook(wb: Workbook) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        result[name] = _clean_headers(header_row or ())
    return result


def _open_workbook(source: str | Path | BinaryIO):
    try:
        if isinstance(source, (str, Path)):
            wb = load_workbook(source, read_only=True, data_only=True)
        else:
            handle_cm = open_workbook_handle(source)
            handle = handle_cm.__enter__()
            try:
                wb = load_workbook(handle, read_only=True, data_only=True)
            except Exception:
                handle_cm.__exit__(None, None, None)
                raise
            wb._claro_handle_cm = handle_cm  # type: ignore[attr-defined]
        return _WorkbookCloser(wb)
    except EpImportError:
        raise
    except Exception as exc:  # noqa: BLE001 — user sees a fixed message; detail stays in the log
        raise EpImportError(title="Could not open EP file", issues=[invalid_workbook()]) from exc


class _WorkbookCloser:
    def __init__(self, workbook: Workbook):
        self._workbook = workbook

    def __enter__(self) -> Workbook:
        return self._workbook

    def __exit__(self, exc_type, exc, tb) -> None:
        handle_cm = getattr(self._workbook, "_claro_handle_cm", None)
        self._workbook.close()
        if handle_cm is not None:
            handle_cm.__exit__(exc_type, exc, tb)


def _json_safe_value(value: Any) -> Any:
    """Coerce openpyxl cell values into JSON-serializable Python types."""
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return int(value)
        return float(value)
    return value


def _clean_headers(header_row: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    for cell in header_row:
        if cell is None:
            headers.append("")
        else:
            headers.append(str(cell).strip())
    while headers and headers[-1] == "":
        headers.pop()
    return headers


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)
