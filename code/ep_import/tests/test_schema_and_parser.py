from pathlib import Path

import pytest
from openpyxl import Workbook

from ep_import.schema import Technology, region_from_filename, resolve_column_name
from ep_import.services.errors import EpImportError
from ep_import.services.importer import import_ep_file
from ep_import.services.layout import validate_headers
from ep_import.services.mapper import site_name

FIXTURES = Path(__file__).parent / "fixtures"


def test_site_name_prefers_single_ran_name():
    assert site_name({"SINGLE RAN NAME": "S01GOGNA83", "ENODEBNAME": "ES01GOGNA83"}, Technology.G4) == "S01GOGNA83"
    assert site_name({"ENODEBNAME": "ES01GOGNA83"}, Technology.G4) == "ES01GOGNA83"
    assert site_name({"SINGLE RAN NAME": "S01", "*BTS NAME": "BS01"}, Technology.G2) == "S01"
    assert site_name({"SINGLE RAN NAME": "S01", "NODEB NAME": "NS01"}, Technology.G3) == "S01"
    assert site_name({"SINGLE RAN NAME": "S01", "ENODEB NAME": "ES01"}, Technology.G5) == "S01"


def test_region_from_filename():
    assert region_from_filename("RNP_SRAN_CO_20260813.xlsx").value == "CO"
    assert region_from_filename("RNP_SRAN_NE_2026.08.13.xlsx").value == "NE"
    assert region_from_filename("RNP_SRAN_BASE_2026.08.10.xlsx").value == "BASE"
    assert region_from_filename("random.xlsx") is None


def test_resolve_column_alias():
    available = {"Region", "STATE", "CELL NAME"}
    assert resolve_column_name("REGION", available) == "Region"
    assert resolve_column_name("STATE", available) == "STATE"
    assert resolve_column_name("gNBId", available) is None


def test_import_minimal_fixture():
    payload = import_ep_file(FIXTURES / "ep_claro_minimal.xlsx")
    assert payload.region == "CO"
    assert payload.counts["2G"] == 2
    assert payload.counts["3G"] == 1
    assert payload.counts["4G"] == 1
    assert payload.counts["5G"] == 2
    assert "RMV_GSM" in payload.ignored_sheets
    sites = payload.sites_for(Technology.G5)
    assert sites == ["ES02ACEPT02"]
    cells = payload.cells_by_site(Technology.G5)["ES02ACEPT02"]
    assert len(cells) == 2
    assert cells[0]["cell_name"] == "52S02ACEPT0201"


def test_import_alias_region_header():
    payload = import_ep_file(FIXTURES / "ep_claro_alias_region.xlsx")
    assert payload.counts["5G"] == 2
    # Region alias should populate region on cells
    assert payload.cells[0].region == "CO"


def test_rejects_missing_sheet(tmp_path):
    path = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "RNP GSM (2G)"
    ws.append(["STATE", "SINGLE RAN NAME", "*BTS NAME", "*GSM CELL NAME", "*LAC", "*CI", "*FREQUENCY OF BCCH", "BSC"])
    ws.append(["AC", "S1", "B1", "C1", 1, 1, 1, "BSC"])
    wb.save(path)

    with pytest.raises(EpImportError) as exc:
        import_ep_file(path)
    codes = {i.code for i in exc.value.issues}
    assert "LAYOUT_MISSING_SHEET" in codes


def test_rejects_missing_column(tmp_path):
    path = tmp_path / "RNP_SRAN_CO_missing_col.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for name in [
        "RNP GSM (2G)",
        "RNP UMTS (3G)",
        "RNP LTE (4G)",
        "RNP NR (5G)",
    ]:
        ws = wb.create_sheet(name)
        if name == "RNP NR (5G)":
            # Missing gNBId
            ws.append(
                [
                    "REGION",
                    "STATE",
                    "SINGLE RAN NAME",
                    "ENODEB NAME",
                    "CELL NAME",
                    "CellId",
                    "FrequencyBand",
                    "PhysicalCellId",
                    "DlNarfcn",
                    "Tracking Area ID",
                ]
            )
            ws.append(["CO", "AC", "S", "E", "C", 1, "N1", 1, 1, 1])
        else:
            # Minimal valid-ish headers for other sheets so only NR fails columns
            headers = {
                "RNP GSM (2G)": [
                    "STATE",
                    "SINGLE RAN NAME",
                    "*BTS NAME",
                    "*GSM CELL NAME",
                    "*LAC",
                    "*CI",
                    "*FREQUENCY OF BCCH",
                    "BSC",
                ],
                "RNP UMTS (3G)": [
                    "REGION",
                    "STATE",
                    "SINGLE RAN NAME",
                    "NODEB NAME",
                    "CELLNAME",
                    "CELL ID",
                    "LAC",
                    "SAC",
                    "RNC ID",
                    "RNC NAME",
                    "PSCRAMBCODE",
                    "UARFCN DOWNLINK",
                ],
                "RNP LTE (4G)": [
                    "REGION",
                    "STATE",
                    "SINGLE RAN NAME",
                    "ENODEBNAME",
                    "CELL NAME",
                    "CELL ID",
                    "ENODEB ID",
                    "TAC",
                    "EARFCN_DL",
                    "PCI",
                ],
            }[name]
            ws.append(headers)
            ws.append(["x"] * len(headers))
    wb.save(path)

    with pytest.raises(EpImportError) as exc:
        import_ep_file(path)
    assert any(i.code == "LAYOUT_MISSING_COLUMN" and i.column == "gNBId" for i in exc.value.issues)


def test_validate_headers_nr_ok():
    headers = [
        "REGION",
        "STATE",
        "SINGLE RAN NAME",
        "ENODEB NAME",
        "CELL NAME",
        "gNBId",
        "CellId",
        "FrequencyBand",
        "PhysicalCellId",
        "DlNarfcn",
        "Tracking Area ID",
    ]
    assert validate_headers("RNP NR (5G)", headers, Technology.G5) == []
